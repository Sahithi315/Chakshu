"""Row-level accuracy of the report extractor against the team's answer key (PDF_RESULT_TRUTH), per layout.
Development layouts: classic, compact, cards. The held-out layout (two_column) runs only with --held-out, once.
python -m src.reports.evaluate [--held-out]"""
import argparse, collections, json
import openpyxl
from .extract import extract_pdf

XLSX, PDFS = "data/report_case_training_dataset.xlsx", "data/synthetic_lab_reports/"
FIELDS = ["test", "value", "unit", "reference_range"]


def same(a, b):
    try:
        return float(a) == float(b)
    except ValueError:
        return str(a).strip() == str(b).strip()


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--held-out", action="store_true"); a = ap.parse_args()
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    layout = {r[1]: r[10] for r in list(wb["PDF_REPORT_INDEX"].iter_rows(values_only=True))[1:]}
    truth = collections.defaultdict(list)
    for r in list(wb["PDF_RESULT_TRUTH"].iter_rows(values_only=True))[1:]:
        truth[r[9]].append(dict(zip(FIELDS, (r[2], r[3], r[4], r[7]))))
    want = {"two_column"} if a.held_out else {"classic", "compact", "cards"}
    stat = {l: collections.Counter() for l in sorted(want)}
    misses = []
    for pdf, l in sorted(layout.items()):
        if l not in want:
            continue
        got = extract_pdf(PDFS + pdf); s = stat[l]; s["pdfs"] += 1
        by_test = {g["test"]: g for g in got}
        s["truth_rows"] += len(truth[pdf]); s["extracted_rows"] += len(got); s["flagged_by_quote_check"] += sum(not g["quote_ok"] for g in got)
        for t in truth[pdf]:
            g = by_test.get(t["test"])
            if g is None:
                s["missed"] += 1; misses.append((pdf, t, None)); continue
            ok = {f: same(g[f], t[f]) for f in FIELDS}
            for f in FIELDS: s[f"{f}_correct"] += ok[f]
            s["row_correct"] += all(ok.values())
            if not all(ok.values()): misses.append((pdf, t, {f: g[f] for f in FIELDS}))
        s["spurious"] += len([g for g in got if g["test"] not in {t["test"] for t in truth[pdf]}])
    print("PDFs were made by the team's own script; accuracy here is not a real-world claim.\n")
    for l, s in stat.items():
        n = s["truth_rows"]
        print(f"{l}: {s['pdfs']} PDFs, {n} answer-key rows, {s['extracted_rows']} extracted | fully correct {s['row_correct']}/{n} | missed {s['missed']} | spurious {s['spurious']} | quote-check flags {s['flagged_by_quote_check']}")
        print("   per field correct:", {f: f"{s[f + '_correct']}/{n}" for f in FIELDS})
    for m in misses[:20]: print("  MISS", m)
    json.dump({l: dict(s) for l, s in stat.items()}, open("outputs/reports_" + ("held_out" if a.held_out else "dev") + ".json", "w"), indent=1)


if __name__ == "__main__":
    main()
